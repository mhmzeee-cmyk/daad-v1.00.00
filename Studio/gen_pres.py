# -*- coding: utf-8 -*-
import json, sys

# Data for all slides
slides_data = json.loads(open(sys.argv[1], encoding='utf-8').read())

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BG=RGBColor(0x0A,0x0A,0x0A); CARD=RGBColor(0x14,0x14,0x14); CARD2=RGBColor(0x1A,0x1A,0x1A)
BORDER=RGBColor(0x28,0x28,0x28); WHITE=RGBColor(0xFF,0xFF,0xFF); GRAY=RGBColor(0xAA,0xAA,0xAA)
GRAY2=RGBColor(0x66,0x66,0x66); GREEN=RGBColor(0x00,0xC9,0xA7); GREEN2=RGBColor(0x00,0xA8,0x8A)
RED=RGBColor(0xE8,0x4D,0x3D); ORANGE=RGBColor(0xF5,0xA6,0x23); BLUE=RGBColor(0x4A,0x9E,0xF5)

def bg(sl):
    f=sl.background.fill; f.solid(); f.fore_color.rgb=BG
def rc(sl,l,t,w,h,c):
    s=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,l,t,w,h); s.fill.solid(); s.fill.fore_color.rgb=c; s.line.fill.background(); return s
def rr(sl,l,t,w,h,c):
    s=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,l,t,w,h); s.fill.solid(); s.fill.fore_color.rgb=c; s.line.color.rgb=BORDER; s.line.width=Pt(0.75); return s
def tx(sl,l,t,w,h,text,sz=18,c=WHITE,b=False,a=PP_ALIGN.RIGHT,f="Segoe UI"):
    tb=sl.shapes.add_textbox(l,t,w,h); tf=tb.text_frame; tf.word_wrap=True; p=tf.paragraphs[0]; p.text=text; p.font.size=Pt(sz); p.font.color.rgb=c; p.font.bold=b; p.font.name=f; p.alignment=a; return tb
def mt(sl,l,t,w,h,lines,sz=15,c=GRAY,a=PP_ALIGN.RIGHT):
    tb=sl.shapes.add_textbox(l,t,w,h); tf=tb.text_frame; tf.word_wrap=True
    for i,ln in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=ln; p.font.size=Pt(sz); p.font.color.rgb=c; p.font.name="Segoe UI"; p.alignment=a; p.space_after=Pt(6)
    return tb
def cd(sl,l,t,w,h,text,sz=12):
    rr(sl,l,t,w,h,CARD); tb=sl.shapes.add_textbox(l+Inches(0.3),t+Inches(0.25),w-Inches(0.6),h-Inches(0.5)); tf=tb.text_frame; tf.word_wrap=True
    for i,ln in enumerate(text.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=ln; p.font.size=Pt(sz); p.font.color.rgb=GREEN; p.font.name="Consolas"; p.alignment=PP_ALIGN.LEFT; p.space_after=Pt(2)
def br(sl): rc(sl,Inches(0),Inches(0),Inches(13.333),Inches(0.04),GREEN)
def ln(sl,l,t,w,c=GREEN): rc(sl,l,t,w,Inches(0.015),c)
def nm(sl,n,t=29): tx(sl,Inches(6),Inches(7.05),Inches(1.3),Inches(0.35),f"{n} / {t}",sz=9,c=GRAY2,a=PP_ALIGN.CENTER,f="Segoe UI Light")
def hd(sl,text,up=Inches(8.5),uw=Inches(4)):
    tx(sl,Inches(0.8),Inches(0.45),Inches(11.7),Inches(0.7),text,sz=32,c=WHITE,b=True,f="Segoe UI Semibold"); ln(sl,up,Inches(1.1),uw)

def I(v): return Inches(v)

prs=Presentation(); prs.slide_width=I(13.333); prs.slide_height=I(7.5)
T=len(slides_data)

for idx, sd in enumerate(slides_data):
    s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    tp=sd["type"]
    
    if tp=="cover":
        rc(s,I(0),I(0),I(13.333),I(0.04),GREEN)
        tx(s,I(1),I(2.5),I(11.3),I(1.2),sd["title"],sz=72,c=WHITE,b=True,a=PP_ALIGN.CENTER,f="Segoe UI Light")
        tx(s,I(1),I(3.8),I(11.3),I(0.6),sd["subtitle"],sz=22,c=GREEN,a=PP_ALIGN.CENTER)
        ln(s,I(5.5),I(4.6),I(2.3))
        for j,line in enumerate(sd.get("lines",[])):
            tx(s,I(1),I(5.0+j*0.6),I(11.3),I(0.5),line["text"],sz=line.get("sz",13),c=GRAY2 if j==0 else GRAY,a=PP_ALIGN.CENTER)
    
    elif tp=="slide":
        br(s)
        h=sd.get("heading",{})
        up=I(h.get("up",8.5)); uw=I(h.get("uw",4))
        hd(s,h["text"],up,uw)
        
        if "items" in sd:
            for i,item in enumerate(sd["items"]):
                y=I(1.6+i*1.35)
                rr(s,I(0.8),y,I(11.7),I(1.1),CARD)
                ln(s,I(11.6),y+I(0.2),I(0.6),GREEN)
                tx(s,I(1.2),y+I(0.1),I(4),I(0.4),item["t"],sz=16,c=GREEN,b=True)
                tx(s,I(1.2),y+I(0.55),I(10.5),I(0.4),item["d"],sz=13,c=GRAY)
        
        if "columns" in sd:
            for i,col in enumerate(sd["columns"]):
                x=I(0.5+i*6.5)
                rr(s,x,I(2.0),I(5.8),I(4.5),CARD)
                tx(s,x+I(0.3),I(2.1),I(5.2),I(0.4),col["title"],sz=16,c=GREEN,b=True)
                ln(s,x+I(0.3),I(2.55),I(4))
                mt(s,x+I(0.3),I(2.7),I(5.2),I(3.5),col["items"],sz=12)
        
        if "table" in sd:
            t=sd["table"]
            hy=I(1.5); rr(s,I(0.8),hy,I(11.7),I(0.5),GREEN)
            for j,h in enumerate(t["headers"]):
                tx(s,I(h["x"]),hy+I(0.05),I(h["w"]),I(0.4),h["text"],sz=13,c=BG,b=True,a=PP_ALIGN.CENTER)
            for i,row in enumerate(t["rows"]):
                y=I(2.15+i*0.58); bgc=CARD if i%2==0 else CARD2
                rr(s,I(0.8),y,I(11.7),I(0.5),bgc)
                for j,cell in enumerate(row):
                    clr=GREEN if j==0 else GRAY
                    tx(s,I(cell["x"]),y+I(0.06),I(cell["w"]),I(0.38),cell["text"],sz=12+int(j==0),c=clr,b=(j==0),a=PP_ALIGN.CENTER)
        
        if "boxes" in sd:
            for box in sd["boxes"]:
                y=I(box["y"])
                rr(s,I(box["x"]),y,I(box["w"]),I(box["h"]),box.get("bg",CARD))
                if box.get("accent"):
                    rc(s,I(box["x"]),y,I(box["w"]),I(0.025),GREEN)
                tx(s,I(box["x"])+I(0.3),y+I(0.2),I(box["w"])-I(0.6),I(0.35),box["title"],sz=16,c=GREEN,b=True)
                tx(s,I(box["x"])+I(0.3),y+I(0.6),I(box["w"])-I(0.6),I(0.4),box["desc"],sz=13,c=GRAY)
        
        if "list" in sd:
            li=sd["list"]
            tx(s,I(li.get("x",0.8)),I(li["y"]),I(li.get("w",5.5)),I(0.4),li.get("label",""),sz=18,c=GREEN,b=True)
            mt(s,I(li.get("x",0.8)),I(li["y"]+0.6),I(li.get("w",5.5)),I(4),li["items"],sz=14)
        
        if "code_block" in sd:
            cb=sd["code_block"]
            cd(s,I(cb["x"]),I(cb["y"]),I(cb["w"]),I(cb["h"]),cb["code"],sz=cb.get("sz",12))
        
        if "flow" in sd:
            fl=sd["flow"]
            for i,step in enumerate(fl["steps"]):
                x=I(fl["sx"])+I(i*fl["gap"])
                rr(s,x,I(fl["y"]),I(fl["bw"]),I(fl["bh"]),GREEN if step.get("highlight") else CARD2)
                tc=BG if step.get("highlight") else WHITE
                tx(s,x,I(fl["y"])+I(0.1),I(fl["bw"]),I(0.4),step["text"],sz=11,c=tc,b=True,a=PP_ALIGN.CENTER)
                if i<len(fl["steps"])-1:
                    tx(s,x+I(fl["bw"])+I(0.05),I(fl["y"])+I(0.08),I(0.2),I(0.4),"\u2192",sz=14,c=GREEN2,a=PP_ALIGN.CENTER)
            if fl.get("note"):
                tx(s,I(0.5),I(fl["note"]["y"]),I(12.3),I(0.5),fl["note"]["text"],sz=13,c=GRAY2,a=PP_ALIGN.CENTER)
        
        if "status_grid" in sd:
            sg=sd["status_grid"]
            xp=[I(0.3),I(3.4),I(6.5),I(9.6)]
            color_map={"green":GREEN,"orange":ORANGE,"blue":BLUE,"red":RED}
            for idx2,(st,items,clr) in enumerate(sg):
                c=color_map.get(clr,GRAY)
                x=xp[idx2]; y=I(1.6)
                rr(s,x,y,I(2.9),I(0.45),c)
                tc=BG if c==GREEN else WHITE
                tx(s,x,y+I(0.05),I(2.9),I(0.35),st,sz=12,c=tc,b=True,a=PP_ALIGN.CENTER)
                for j,item in enumerate(items):
                    iy=y+I(0.6+j*0.58)
                    tx(s,x+I(0.1),iy,I(2.7),I(0.5),"\u2022 "+item,sz=10,c=GRAY,a=PP_ALIGN.RIGHT)
        
        if "note" in sd:
            n=sd["note"]
            tx(s,I(n.get("x",0.5)),I(n["y"]),I(n.get("w",12.3)),I(0.5),n["text"],sz=n.get("sz",11),c=n.get("c",GRAY2),a=PP_ALIGN.CENTER)
    
    elif tp=="conclusion":
        rc(s,I(0),I(0),I(13.333),I(0.04),GREEN)
        tx(s,I(1),I(1.5),I(11.3),I(1),sd["title"],sz=48,c=WHITE,b=True,a=PP_ALIGN.CENTER,f="Segoe UI Light")
        for i,item in enumerate(sd["items"]):
            y=I(2.5+i*1.0)
            rr(s,I(1.5),y,I(10.3),I(0.8),CARD)
            rc(s,I(1.5),y,I(0.06),I(0.8),GREEN)
            tx(s,I(1.8),y+I(0.15),I(9.7),I(0.5),item,sz=18,c=WHITE,a=PP_ALIGN.RIGHT)
        tx(s,I(1),I(6.2),I(11.3),I(0.8),sd["quote"],sz=22,c=GREEN,b=True,a=PP_ALIGN.CENTER)
    
    elif tp=="sources":
        br(s)
        hd(s,sd["heading"],I(7.5),I(5))
        mt(s,I(0.5),I(1.8),I(12.3),I(5),sd["items"],sz=14)
    
    nm(s,idx+1,T)

output=slides_data[0].get("output","DAAD_LANGUAGE_PRESENTATION.pptx")
prs.save(output)
print(f"Saved: {output}")
print(f"Slides: {len(prs.slides)}")
