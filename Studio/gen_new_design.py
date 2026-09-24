# -*- coding: utf-8 -*-
"""
================================================================================
DAAD PRESENTATION ENGINE - SYSTEM ENTRY POINT & COMPILER
================================================================================
Imports all modular components to build the professional 29-slide presentation.
================================================================================
"""

import os
import json
from pptx import Presentation
from daad_presenter.theme import Dimensions
from daad_presenter.shapes import ShapeService
from daad_presenter.renderers_part1 import RenderersPart1
from daad_presenter.renderers_part2 import RenderersPart2
from daad_presenter.renderers_part3 import RenderersPart3

def main():
    print("[INIT] Loading Presentation JSON data schema...")
    with open('slides_data.json', 'r', encoding='utf-8') as f:
        slides_data = json.load(f)
        
    prs = Presentation()
    prs.slide_width = Dimensions.SCREEN_W
    prs.slide_height = Dimensions.SCREEN_H
    
    print("[COMPILE] Directing slide data down the specialized layout pipelines...")
    for idx, sd in enumerate(slides_data):
        slide_idx = idx + 1
        
        # Heading safe loading logic
        heading_data = sd.get("heading", {})
        if isinstance(heading_data, dict):
            slide_title = heading_data.get("text", sd.get("title", ""))
        else:
            slide_title = heading_data if heading_data else sd.get("title", "")
            
        print(f" -> Compiling Slide {slide_idx:02d}: {slide_title}")
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        
        # Advanced slide-by-slide layout routing
        if sd.get("type") == "cover":
            RenderersPart1.render_cover(slide, sd, slide_idx)
        elif slide_idx == 2:
            RenderersPart1.render_slide_2_what_is_daad(slide, sd, slide_idx)
        elif slide_idx == 3:
            RenderersPart1.render_slide_3_core_components(slide, sd, slide_idx)
        elif slide_idx == 4:
            RenderersPart1.render_slide_4_audience(slide, sd, slide_idx)
        elif slide_idx == 5:
            RenderersPart1.render_slide_5_philosophy(slide, sd, slide_idx)
        elif slide_idx == 6:
            RenderersPart1.render_slide_6_split_screen(slide, sd, slide_idx)
        elif slide_idx == 7:
            RenderersPart1.render_slide_7_quick_look(slide, sd, slide_idx)
        elif slide_idx == 8:
            RenderersPart1.render_slide_8_components_matrix(slide, sd, slide_idx)
        elif slide_idx == 9:
            RenderersPart2.render_slide_9_flowchart(slide, sd, slide_idx)
        elif slide_idx == 10:
            RenderersPart1.render_slide_10_lexer(slide, sd, slide_idx)
        elif slide_idx == 11:
            RenderersPart2.render_slide_11_parser(slide, sd, slide_idx)
        elif slide_idx == 12:
            RenderersPart2.render_slide_12_ast(slide, sd, slide_idx)
        elif slide_idx == 13:
            RenderersPart2.render_slide_13_type_system(slide, sd, slide_idx)
        elif slide_idx == 14:
            RenderersPart2.render_slide_14_stdlib(slide, sd, slide_idx)
        elif slide_idx == 15:
            RenderersPart2.render_slide_15_status(slide, sd, slide_idx)
        elif slide_idx == 16:
            RenderersPart2.render_slide_16_achievements(slide, sd, slide_idx)
        elif slide_idx == 17:
            RenderersPart2.render_slide_17_incomplete(slide, sd, slide_idx)
        elif slide_idx == 18:
            RenderersPart2.render_slide_18_now(slide, sd, slide_idx)
        elif slide_idx == 19:
            RenderersPart2.render_slide_19_gap_analysis(slide, sd, slide_idx)
        elif slide_idx == 20:
            RenderersPart2.render_slide_20_roadmap_full(slide, sd, slide_idx)
        elif slide_idx == 21:
            RenderersPart3.render_slide_21_ir_why(slide, sd, slide_idx)
        elif slide_idx == 22:
            RenderersPart3.render_slide_22_ir_conversion(slide, sd, slide_idx)
        elif slide_idx == 23:
            RenderersPart3.render_slide_23_processors(slide, sd, slide_idx)
        elif slide_idx == 24:
            RenderersPart3.render_slide_24_roadmap(slide, sd, slide_idx)
        elif slide_idx == 25:
            RenderersPart3.render_slide_25_strengths(slide, sd, slide_idx)
        elif slide_idx == 26:
            RenderersPart3.render_slide_26_challenges(slide, sd, slide_idx)
        elif slide_idx == 27:
            RenderersPart3.render_slide_27_future_vision(slide, sd, slide_idx)
        elif slide_idx == 28:
            RenderersPart3.render_slide_28_conclusion(slide, sd, slide_idx)
        elif slide_idx == 29:
            RenderersPart3.render_slide_29_sources(slide, sd, slide_idx)
        else:
            RenderersPart3.render_generic_grid(slide, sd, slide_idx)
            
        # Add ambient vector decorations to every slide (corner motifs, edge nodes, bottom data band)
        ShapeService.decorate_slide(slide, slide_idx)
            
    # Save the ultimate compiled asset
    out_filename = "DAAD_LANGUAGE_PRESENTATION.pptx"
    prs.save(out_filename)
    print(f"\n[SUCCESS] Compiled Presentation Saved: {out_filename}")
    print(f" -> Verification Passed: {len(prs.slides)} Slides compiled successfully.")

if __name__ == "__main__":
    main()
